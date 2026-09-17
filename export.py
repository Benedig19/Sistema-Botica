"""
export.py - Exportación profesional a Excel y PDF
"""

import sqlite3
from datetime import datetime
from pathlib import Path

# Excel
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_DISPONIBLE = True
except ImportError:
    EXCEL_DISPONIBLE = False

# PDF
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    PDF_DISPONIBLE = True
except ImportError:
    PDF_DISPONIBLE = False

DB_PATH = Path("botica.db")

def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def exportar_inventario_excel(ruta_salida):
    """Exporta todo el inventario activo a Excel profesional."""
    if not EXCEL_DISPONIBLE:
        raise ImportError("Instale openpyxl: pip install openpyxl")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM productos WHERE activo = 1 ORDER BY nombre")
    productos = cursor.fetchall()
    conn.close()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="1976d2", end_color="1976d2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Encabezados
    headers = ['ID', 'Código', 'Nombre', 'Descripción', 'Categoría', 'P. Compra', 'P. Venta', 
               'Stock', 'Stock Mín.', 'Vencimiento', 'Laboratorio']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    
    # Datos
    for row, p in enumerate(productos, 2):
        ws.cell(row=row, column=1, value=p['id']).border = thin_border
        ws.cell(row=row, column=2, value=p['codigo']).border = thin_border
        ws.cell(row=row, column=3, value=p['nombre']).border = thin_border
        ws.cell(row=row, column=4, value=p['descripcion']).border = thin_border
        ws.cell(row=row, column=5, value=p['categoria']).border = thin_border
        ws.cell(row=row, column=6, value=p['precio_compra']).border = thin_border
        ws.cell(row=row, column=7, value=p['precio_venta']).border = thin_border
        ws.cell(row=row, column=8, value=p['stock']).border = thin_border
        ws.cell(row=row, column=9, value=p['stock_minimo']).border = thin_border
        ws.cell(row=row, column=10, value=p['fecha_vencimiento']).border = thin_border
        ws.cell(row=row, column=11, value=p['laboratorio']).border = thin_border
        
        # Color rojo si stock bajo
        if p['stock'] <= p['stock_minimo']:
            for col in range(1, 12):
                ws.cell(row=row, column=col).fill = PatternFill(start_color="ffcccc", end_color="ffcccc", fill_type="solid")
    
    # Ajustar anchos
    anchos = [6, 12, 25, 30, 15, 12, 12, 10, 12, 14, 20]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    
    # Congelar paneles
    ws.freeze_panes = 'A2'
    
    wb.save(ruta_salida)
    return ruta_salida

def exportar_ventas_excel(fecha_inicio, fecha_fin, ruta_salida):
    """Exporta ventas del período a Excel."""
    if not EXCEL_DISPONIBLE:
        raise ImportError("Instale openpyxl: pip install openpyxl")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT v.*, c.nombre as cliente_nombre 
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        WHERE DATE(v.fecha) BETWEEN ? AND ?
        ORDER BY v.fecha DESC
    ''', (fecha_inicio, fecha_fin))
    ventas = cursor.fetchall()
    conn.close()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2e7d32", end_color="2e7d32", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    headers = ['ID', 'Fecha', 'Cliente', 'Total', 'Método Pago', 'Usuario']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    
    total = 0
    for row, v in enumerate(ventas, 2):
        cliente = v['cliente_nombre'] or "General"
        ws.cell(row=row, column=1, value=v['id']).border = thin_border
        ws.cell(row=row, column=2, value=v['fecha']).border = thin_border
        ws.cell(row=row, column=3, value=cliente).border = thin_border
        ws.cell(row=row, column=4, value=v['total']).border = thin_border
        ws.cell(row=row, column=5, value=v['metodo_pago']).border = thin_border
        ws.cell(row=row, column=6, value=v['usuario'] or 'Admin').border = thin_border
        total += v['total']
    
    # Fila de totales
    row_total = len(ventas) + 2
    ws.cell(row=row_total, column=3, value="TOTAL:").font = Font(bold=True)
    ws.cell(row=row_total, column=4, value=total).font = Font(bold=True, size=12)
    
    anchos = [8, 20, 25, 12, 15, 15]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    
    ws.freeze_panes = 'A2'
    wb.save(ruta_salida)
    return ruta_salida

def exportar_ventas_pdf(fecha_inicio, fecha_fin, ruta_salida):
    """Exporta reporte de ventas a PDF profesional."""
    if not PDF_DISPONIBLE:
        raise ImportError("Instale reportlab: pip install reportlab")
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT v.*, c.nombre as cliente_nombre 
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        WHERE DATE(v.fecha) BETWEEN ? AND ?
        ORDER BY v.fecha DESC
    ''', (fecha_inicio, fecha_fin))
    ventas = cursor.fetchall()
    conn.close()
    
    doc = SimpleDocTemplate(ruta_salida, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=30)
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1976d2'),
        alignment=TA_CENTER,
        spaceAfter=20
    )
    
    elements.append(Paragraph("💊 BOTICA - REPORTE DE VENTAS", title_style))
    elements.append(Paragraph(f"Período: {fecha_inicio} al {fecha_fin}", styles['Normal']))
    elements.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 20))
    
    # Tabla
    data = [['ID', 'Fecha', 'Cliente', 'Total', 'Método', 'Usuario']]
    total = 0
    for v in ventas:
        cliente = v['cliente_nombre'] or "General"
        data.append([
            str(v['id']),
            v['fecha'][:16],
            cliente[:25],
            f"S/ {v['total']:.2f}",
            v['metodo_pago'],
            v['usuario'] or 'Admin'
        ])
        total += v['total']
    
    data.append(['', '', 'TOTAL:', f"S/ {total:.2f}", '', ''])
    
    table = Table(data, colWidths=[40, 100, 130, 70, 70, 70])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976d2')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -2), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -2), colors.HexColor('#f5f5f5')),
        ('GRID', (0, 0), (-1, -2), 1, colors.grey),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e3f2fd')),
    ]))
    
    elements.append(table)
    elements.append(Spacer(1, 30))
    elements.append(Paragraph(f"Total de ventas en el período: <b>S/ {total:.2f}</b>", styles['Normal']))
    
    doc.build(elements)
    return ruta_salida